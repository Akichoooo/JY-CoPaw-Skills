#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""预警准确率日报 / 月报生成器。

日报:
    python accuracy_report.py daily <导出.xlsx> --date 2026-10-26 \
        --false "夜间区域入侵=3" --false "区域入侵=1" --false "烟雾识别=all" \
        --reason "夜间区域入侵=算法误识别：把门识别成人" \
        --reason "区域入侵=算法误识别：022镜头把围栏识别成人"

月报:
    python accuracy_report.py monthly --month 2026-10 --ledger <目录>/台账/准确率台账.csv

默认输出：准确率日报_YYYY-MM-DD.xlsx（每天一份，不覆盖昨天）、
准确率月报_YYYY-MM.xlsx；台账收在 台账/准确率台账.csv，供跨天汇总。
"""

from __future__ import annotations

import argparse
import calendar
import csv
import filecmp
import re
import shutil
import sys
import traceback
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.marker import Marker
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    print("错误：缺少 openpyxl，请先运行 pip install openpyxl")
    sys.exit(1)


FONT_NAME = "微软雅黑"
COLOR_NAVY = "1F4E78"
COLOR_NAVY_LIGHT = "DDEBF7"
COLOR_GREY = "595959"
COLOR_BORDER = "BFBFBF"
COLOR_BAD_FILL = "FCE4E4"
COLOR_BAD_FONT = "C00000"
COLOR_WARN_FILL = "FFF2CC"
COLOR_WARN_FONT = "9C6500"
COLOR_OK_FILL = "E2EFDA"
COLOR_OK_FONT = "375623"
COLOR_MISSING_FILL = "F2F2F2"

LEDGER_COLUMNS = [
    "报表日期",
    "预警技能",
    "预警总数",
    "误报数",
    "误报原因",
    "准确数",
    "准确率",
    "数据源",
    "统计开始",
    "统计结束",
    "生成时间",
]


class UserError(Exception):
    """可直接展示给用户的错误。"""


# ---------------------------------------------------------------------------
# 通用小工具
# ---------------------------------------------------------------------------

def _init_stdout() -> None:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8", errors="replace")
            except Exception:
                pass


def parse_dt(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time())
    if value is None:
        return None
    text = str(value).strip()
    match = re.match(r"(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}:\d{2})", text)
    if not match:
        return None
    try:
        return datetime.strptime(match.group(1) + " " + match.group(2), "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None


def fmt_dt(value) -> str:
    dt = parse_dt(value)
    if dt:
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    if value is None or str(value).strip() == "":
        return "—"
    return str(value).strip()[:19]


def thin_border() -> Border:
    side = Side(style="thin", color=COLOR_BORDER)
    return Border(left=side, right=side, top=side, bottom=side)


def add_title(ws, text: str, ncols: int, row: int = 1, size: int = 14) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row, 1, text)
    cell.font = Font(name=FONT_NAME, size=size, bold=True, color=COLOR_NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 28


def add_note(ws, row: int, text: str, ncols: int, fill: str | None = None,
             font_color: str = COLOR_GREY, bold: bool = False) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row, 1, text)
    cell.font = Font(name=FONT_NAME, size=9, color=font_color, bold=bold)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if fill:
        for col in range(1, ncols + 1):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=fill)
            ws.cell(row, col).border = thin_border()
    ws.row_dimensions[row].height = 22 if len(text) <= 70 else 34


def add_section(ws, row: int, text: str, ncols: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row, 1, text)
    cell.font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24


def write_header_row(ws, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row, col, text)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=COLOR_NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border()
    ws.row_dimensions[row].height = 26


def write_row(ws, row: int, values: list, percent_cols=(), bold: bool = False,
              fill: str | None = None, font_color: str | None = None,
              center_first: bool = False, left_cols=()) -> None:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.font = Font(
            name=FONT_NAME, size=10, bold=bold,
            color=font_color if font_color else "000000",
        )
        cell.border = thin_border()
        centered = (center_first or col > 1) and col not in left_cols
        cell.alignment = Alignment(
            horizontal="center" if centered else "left",
            vertical="center", wrap_text=True,
        )
        if col in percent_cols:
            cell.number_format = "0.0%"
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)


def safe_save(wb, path: Path) -> None:
    try:
        wb.save(path)
    except PermissionError as exc:
        raise UserError(
            f"无法写入 {path}，文件可能正在 Excel 中打开。请关闭后重试。原始错误：{exc}"
        )
    except OSError as exc:
        raise UserError(f"无法写入 {path}：{exc}")


DAILY_RE = re.compile(r"^准确率日报_(\d{4}-\d{2}-\d{2})\.xlsx$")
MONTHLY_RE = re.compile(r"^准确率月报_(\d{4}-\d{2})\.xlsx$")


def _move_report(src: Path, dest: Path) -> str:
    try:
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists():
            if dest.is_file() and filecmp.cmp(src, dest, shallow=False):
                src.unlink()
                return "skipped"
            return "conflict"
        shutil.move(str(src), str(dest))
        return "moved"
    except OSError:
        return "locked"


def tidy_reports(outdir: Path) -> tuple[int, int, list[str]]:
    """根目录只留最新一份日报和月报，其余收进 归档/YYYY-MM/。"""
    dailies = []
    for path in outdir.glob("准确率日报_*.xlsx"):
        match = DAILY_RE.match(path.name)
        if match:
            dailies.append((match.group(1), path))
    monthlies = []
    for path in outdir.glob("准确率月报_*.xlsx"):
        match = MONTHLY_RE.match(path.name)
        if match:
            monthlies.append((match.group(1), path))

    if not dailies and not monthlies:
        return 0, 0, []

    latest_daily = max(date for date, _ in dailies) if dailies else None
    latest_monthly = max(month for month, _ in monthlies) if monthlies else None
    moved = 0
    skipped = 0
    problems = []

    for date, path in sorted(dailies):
        if date == latest_daily:
            continue
        dest = outdir / "归档" / date[:7] / "日报" / path.name
        result = _move_report(path, dest)
        if result == "moved":
            moved += 1
        elif result == "skipped":
            skipped += 1
        else:
            reason = "目标已存在且内容不同" if result == "conflict" else "文件被占用"
            problems.append(f"{path.name}（{reason}）")

    for month, path in sorted(monthlies):
        if month == latest_monthly:
            continue
        dest = outdir / "归档" / month / path.name
        result = _move_report(path, dest)
        if result == "moved":
            moved += 1
        elif result == "skipped":
            skipped += 1
        else:
            reason = "目标已存在且内容不同" if result == "conflict" else "文件被占用"
            problems.append(f"{path.name}（{reason}）")

    if moved or skipped:
        extra = f"，跳过 {skipped} 个已存在" if skipped else ""
        print(f"整理：归档 {moved} 个旧报告{extra}")
    for problem in problems:
        print(f"  未归档：{problem}")
    return moved, skipped, problems


# ---------------------------------------------------------------------------
# 读取预警导出
# ---------------------------------------------------------------------------

def find_header(ws):
    max_scan = min(ws.max_row, 30)
    for row in ws.iter_rows(min_row=1, max_row=max_scan):
        values = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in row
        ]
        if "预警技能" in values:
            headers = {value: idx for idx, value in enumerate(values) if value}
            return row[0].row, headers
    return None, None


def read_export(path: Path, sheet_name: str | None = None) -> dict:
    if not path.exists():
        raise UserError(f"找不到导出文件：{path}")
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise UserError(f"只支持 .xlsx / .xlsm，收到：{path.suffix}。.xls 请先另存为 .xlsx。")
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as exc:
        raise UserError(f"无法读取 Excel：{exc}")

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise UserError(
                f"工作表“{sheet_name}”不存在。现有工作表：{'、'.join(wb.sheetnames)}"
            )
        candidates = [wb[sheet_name]]
    else:
        candidates = list(wb.worksheets)

    found = None
    for ws in candidates:
        header_row, headers = find_header(ws)
        if header_row:
            found = (ws, header_row, headers)
            break
    if not found:
        raise UserError(
            "没有找到包含“预警技能”的表头。请确认这是预警平台的“预警详情”导出表。"
        )

    ws, header_row, headers = found
    skill_idx = headers["预警技能"]
    time_idx = headers.get("预警时间")
    id_idx = headers.get("预警ID")
    point_idx = headers.get("预警点位")
    org_idx = headers.get("所属组织")

    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        def get(idx):
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        skill = get(skill_idx)
        skill = str(skill).strip() if skill is not None else ""
        if not skill:
            continue
        records.append({
            "skill": skill,
            "time": get(time_idx),
            "id": get(id_idx),
            "point": get(point_idx),
            "org": get(org_idx),
        })

    if not records:
        raise UserError("表头下没有有效预警记录。")

    counts = Counter(item["skill"] for item in records)
    times = [parse_dt(item["time"]) for item in records]
    times = [item for item in times if item]
    ids = [
        str(item["id"]).strip()
        for item in records
        if item["id"] not in (None, "")
    ]
    duplicates = [key for key, value in Counter(ids).items() if value > 1]

    return {
        "source": path.name,
        "sheet": ws.title,
        "counts": counts,
        "total": len(records),
        "records": records,
        "time_min_dt": min(times) if times else None,
        "time_max_dt": max(times) if times else None,
        "time_min": fmt_dt(min(times)) if times else "—",
        "time_max": fmt_dt(max(times)) if times else "—",
        "duplicates": duplicates,
    }


def exclude_skills(data: dict, excluded: list[str]) -> dict:
    names = [name.strip() for name in excluded if name.strip()]
    if not names:
        return data

    known = set(data["counts"])
    unknown = [name for name in names if name not in known]
    if unknown:
        available = "、".join(sorted(known, key=lambda key: (-data["counts"][key], key)))
        raise UserError(
            f"要排除的技能不在导出表中：{'、'.join(unknown)}。可选技能：{available}"
        )

    excluded_set = set(names)
    excluded_counts = {name: data["counts"][name] for name in excluded_set}
    records = [item for item in data["records"] if item["skill"] not in excluded_set]
    if not records:
        raise UserError("排除后没有剩余预警记录，无法生成日报。")

    counts = Counter(item["skill"] for item in records)
    times = [parse_dt(item["time"]) for item in records]
    times = [item for item in times if item]
    ids = [
        str(item["id"]).strip()
        for item in records
        if item["id"] not in (None, "")
    ]
    result = dict(data)
    result.update({
        "counts": counts,
        "total": len(records),
        "records": records,
        "time_min_dt": min(times) if times else None,
        "time_max_dt": max(times) if times else None,
        "time_min": fmt_dt(min(times)) if times else "—",
        "time_max": fmt_dt(max(times)) if times else "—",
        "duplicates": [key for key, value in Counter(ids).items() if value > 1],
        "excluded_counts": excluded_counts,
    })
    return result


# ---------------------------------------------------------------------------
# 日报
# ---------------------------------------------------------------------------

def parse_false_specs(specs: list[str], counts: Counter):
    false = {skill: 0 for skill in counts}
    specified: dict[str, int] = {}
    all_skills: set[str] = set()

    for spec in specs:
        if "=" not in spec:
            raise UserError(f"误报参数格式应为 技能=数量 或 技能=all，收到：{spec}")
        name, raw = spec.split("=", 1)
        name = name.strip()
        raw = raw.strip()
        if name not in counts:
            available = "、".join(sorted(counts, key=lambda key: (-counts[key], key)))
            raise UserError(f"技能“{name}”不在导出表中。可选技能：{available}")
        if name in specified:
            raise UserError(f"技能“{name}”的误报数重复指定。")
        if raw.lower() in ("all", "全部", "全", "都", "全误报"):
            number = counts[name]
            all_skills.add(name)
        else:
            try:
                number = int(raw)
            except ValueError:
                raise UserError(
                    f"技能“{name}”的误报数“{raw}”不是整数，也不是 all。"
                )
            if number < 0:
                raise UserError(f"技能“{name}”的误报数不能为负数。")
            if number > counts[name]:
                raise UserError(
                    f"技能“{name}”误报 {number} 条，超过总数 {counts[name]} 条。"
                )
        specified[name] = number
        false[name] = number

    assumed = [skill for skill in counts if skill not in specified]
    return false, specified, all_skills, assumed


def parse_reason_specs(specs: list[str], counts: Counter, false: dict) -> dict[str, str]:
    reasons: dict[str, str] = {}
    if not specs:
        return reasons
    if not any(false.values()):
        raise UserError(
            "提供了 --reason，但所有技能的误报数都是 0；请先用 --false 指定误报条数。"
        )

    for spec in specs:
        if "=" not in spec:
            raise UserError(f"误报原因格式应为 技能=原因，收到：{spec}")
        name, raw = spec.split("=", 1)
        name = name.strip()
        reason = re.sub(r"\s+", " ", raw).strip()
        if name not in counts:
            available = "、".join(sorted(counts, key=lambda key: (-counts[key], key)))
            raise UserError(f"技能“{name}”不在导出表中。可选技能：{available}")
        if name in reasons:
            raise UserError(f"技能“{name}”的误报原因重复指定。")
        if not reason:
            raise UserError(f"技能“{name}”的误报原因为空。")
        if false.get(name, 0) <= 0:
            raise UserError(
                f"技能“{name}”的误报数为 0，不需要填写误报原因。"
            )
        reasons[name] = reason

    return reasons


def build_daily_workbook(report_date: str, data: dict, false: dict,
                         specified: dict, all_skills: set,
                         assumed: list, reasons: dict) -> tuple[Workbook, list[dict]]:
    total = data["total"]
    false_total = sum(false.values())
    accurate_total = total - false_total
    accuracy_total = accurate_total / total
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    excluded_counts = data.get("excluded_counts") or {}
    excluded_total = sum(excluded_counts.values())

    skills = sorted(data["counts"], key=lambda key: (-data["counts"][key], key))
    records = []
    for skill in skills:
        count = data["counts"][skill]
        mis = false[skill]
        records.append({
            "报表日期": report_date,
            "预警技能": skill,
            "预警总数": count,
            "误报数": mis,
            "误报原因": reasons.get(skill, ""),
            "准确数": count - mis,
            "准确率": (count - mis) / count,
            "数据源": data["source"],
            "统计开始": data["time_min"],
            "统计结束": data["time_max"],
            "生成时间": generated_at,
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "准确率统计"
    ws.sheet_view.showGridLines = False
    ncols = 7

    add_title(ws, "预警准确率日报", ncols)
    total_text = f"预警总数：{total} 条"
    if excluded_counts:
        total_text += f"（已排除 {excluded_total} 条测试数据）"
    add_note(
        ws, 2,
        f"报表日期：{report_date}　｜　统计区间：{data['time_min']} 至 {data['time_max']}"
        f"　｜　数据源：{data['source']}　｜　{total_text}",
        ncols,
    )
    add_note(
        ws, 3,
        "口径：准确率 =（预警总数 − 误报数）÷ 预警总数；误报数与误报原因均为人工核对结果。",
        ncols,
    )

    header_row = 5
    write_header_row(ws, header_row, [
        "预警技能", "预警总数", "误报数", "准确数", "准确率", "误报原因", "核对方式",
    ])

    row = header_row + 1
    for skill in skills:
        count = data["counts"][skill]
        mis = false[skill]
        accuracy = (count - mis) / count
        if skill in all_skills:
            note = "全部误报"
        elif skill in specified:
            note = f"指定 {mis} 条"
        else:
            note = "未提及，按 0 误报"
        low = accuracy < 0.9
        reason_text = reasons.get(skill, "") or ("—" if mis == 0 else "未填写")
        write_row(
            ws, row,
            [skill, count, mis, count - mis, accuracy, reason_text, note],
            percent_cols=(5,),
            fill=COLOR_BAD_FILL if low else None,
            font_color=COLOR_BAD_FONT if low else None,
            bold=low,
            left_cols=(6,),
        )
        ws.row_dimensions[row].height = 28 if len(reason_text) > 18 else 22
        row += 1

    write_row(
        ws, row, ["合计", total, false_total, accurate_total, accuracy_total, "", ""],
        percent_cols=(5,), bold=True, fill=COLOR_NAVY_LIGHT,
    )
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.border = Border(
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
            top=Side(style="double", color=COLOR_NAVY),
            bottom=Side(style="thin", color=COLOR_BORDER),
        )
    ws.row_dimensions[row].height = 24

    note_row = row + 2
    specified_parts = []
    for skill in skills:
        if skill not in specified:
            continue
        if skill in all_skills:
            part = f"{skill} {false[skill]} 条（全部误报）"
        else:
            part = f"{skill} {false[skill]} 条"
        if reasons.get(skill):
            part += f"（{reasons[skill]}）"
        specified_parts.append(part)
    if specified_parts:
        add_note(ws, note_row, "误报来源：" + "；".join(specified_parts) + "。", ncols)
        note_row += 1
    missing_reason = [
        f"{skill}（{false[skill]} 条）"
        for skill in skills
        if false[skill] > 0 and not reasons.get(skill)
    ]
    if missing_reason:
        add_note(
            ws, note_row,
            "以下误报未填写原因：" + "、".join(missing_reason)
            + "。可补充 --reason 后重跑。",
            ncols, fill=COLOR_WARN_FILL, font_color=COLOR_WARN_FONT, bold=True,
        )
        note_row += 1
    if excluded_counts:
        add_note(
            ws, note_row,
            "已排除、不计入统计："
            + "、".join(
                f"{skill} {count} 条"
                for skill, count in sorted(excluded_counts.items())
            )
            + "。",
            ncols, fill=COLOR_MISSING_FILL, font_color=COLOR_GREY,
        )
        note_row += 1
    if assumed:
        add_note(
            ws, note_row,
            "未提及、按 0 误报的技能：" + "、".join(assumed)
            + "。若其中也有误报，请补充 --false 后重跑。",
            ncols, fill=COLOR_WARN_FILL, font_color=COLOR_WARN_FONT, bold=True,
        )
        note_row += 1
    if data["duplicates"]:
        add_note(
            ws, note_row,
            f"注意：预警ID有重复（{len(data['duplicates'])} 个），本表仍按条数统计，请核对导出数据。",
            ncols, fill=COLOR_WARN_FILL, font_color=COLOR_WARN_FONT,
        )
        note_row += 1
    add_note(ws, note_row, f"生成时间：{generated_at}", ncols)

    widths = {"A": 20, "B": 11, "C": 10, "D": 10, "E": 11, "F": 40, "G": 18}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A6"

    # 明细核对
    ws2 = wb.create_sheet("明细核对")
    ws2.sheet_view.showGridLines = False
    add_title(ws2, "预警准确率明细核对", 7, size=13)
    add_note(
        ws2, 2,
        f"报表日期：{report_date}　｜　数据源：{data['source']}　｜　{total_text}",
        7,
    )
    write_header_row(ws2, 4, [
        "预警技能", "预警总数", "误报数", "准确数", "准确率", "误报原因", "核对说明",
    ])
    row2 = 5
    for skill in skills:
        count = data["counts"][skill]
        mis = false[skill]
        accuracy = (count - mis) / count
        if skill in all_skills:
            note = "人工核对：全部误报"
        elif skill in specified:
            note = f"人工核对：{mis} 条误报"
        else:
            note = "未提及，按 0 误报"
        reason_text = reasons.get(skill, "") or ("—" if mis == 0 else "未填写")
        write_row(
            ws2, row2,
            [skill, count, mis, count - mis, accuracy, reason_text, note],
            percent_cols=(5,),
            font_color=COLOR_BAD_FONT if accuracy < 0.9 else None,
            bold=accuracy < 0.9,
            left_cols=(6,),
        )
        ws2.row_dimensions[row2].height = 28 if len(reason_text) > 18 else 22
        row2 += 1
    note2_row = row2 + 1
    if excluded_counts:
        add_note(
            ws2, note2_row,
            "已排除、不计入统计："
            + "、".join(
                f"{skill} {count} 条"
                for skill, count in sorted(excluded_counts.items())
            )
            + "。",
            7, fill=COLOR_MISSING_FILL,
        )
        note2_row += 1
    add_note(
        ws2, note2_row,
        "“未提及，按 0 误报”表示当天用户只报了其他技能的误报数；如该技能实际有误报，请补充后重跑。",
        7,
    )
    for col, width in {
        "A": 20, "B": 11, "C": 10, "D": 10, "E": 11, "F": 40, "G": 24,
    }.items():
        ws2.column_dimensions[col].width = width
    ws2.freeze_panes = "A5"

    # 数据记录（机器可读，方便台账丢失时恢复）
    ws3 = wb.create_sheet("数据记录")
    ws3.sheet_view.showGridLines = False
    write_header_row(ws3, 1, LEDGER_COLUMNS)
    row3 = 2
    for item in records:
        write_row(
            ws3, row3,
            [
                item["报表日期"], item["预警技能"], item["预警总数"],
                item["误报数"], item["误报原因"], item["准确数"], item["准确率"],
                item["数据源"], item["统计开始"], item["统计结束"],
                item["生成时间"],
            ],
            percent_cols=(7,),
            left_cols=(5,),
        )
        row3 += 1
    widths3 = {
        "A": 12, "B": 18, "C": 10, "D": 10, "E": 40,
        "F": 10, "G": 10, "H": 26, "I": 20, "J": 20, "K": 20,
    }
    for col, width in widths3.items():
        ws3.column_dimensions[col].width = width
    ws3.freeze_panes = "A2"

    return wb, records


def update_ledger(ledger_path: Path, report_date: str, records: list[dict]) -> None:
    existing = []
    if ledger_path.exists():
        with open(ledger_path, "r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            fields = reader.fieldnames or []
            if "报表日期" not in fields or "预警技能" not in fields:
                raise UserError(
                    f"台账 {ledger_path} 缺少“报表日期”或“预警技能”列，请检查。"
                )
            for item in reader:
                if str(item.get("报表日期", "")).strip() == report_date:
                    continue
                existing.append(item)

    merged = existing + [
        {column: item.get(column, "") for column in LEDGER_COLUMNS}
        for item in records
    ]
    merged.sort(key=lambda item: (
        str(item.get("报表日期", "")),
        str(item.get("预警技能", "")),
    ))

    ledger_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with open(ledger_path, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=LEDGER_COLUMNS)
            writer.writeheader()
            for item in merged:
                row = {column: item.get(column, "") for column in LEDGER_COLUMNS}
                # 准确率统一写成小数，月报按数值解析
                try:
                    row["准确率"] = f"{float(row['准确率']):.6f}"
                except (TypeError, ValueError):
                    pass
                writer.writerow(row)
    except PermissionError as exc:
        raise UserError(
            f"无法写入台账 {ledger_path}，文件可能正在 Excel 中打开。请关闭后重试。原始错误：{exc}"
        )


def cmd_daily(args) -> None:
    export = Path(args.export).expanduser().resolve()
    data = read_export(export, args.sheet)
    data = exclude_skills(data, args.exclude)

    if args.list_skills:
        if data.get("excluded_counts"):
            print(
                "已排除、不统计："
                + "、".join(
                    f"{skill} {count} 条"
                    for skill, count in sorted(data["excluded_counts"].items())
                )
            )
        print(f"技能清单（{data['total']} 条预警）：")
        for skill in sorted(data["counts"], key=lambda key: (-data["counts"][key], key)):
            print(f"  {skill}: {data['counts'][skill]} 条")
        return

    if args.date:
        report_date = args.date.strip()
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", report_date):
            raise UserError(f"--date 格式应为 YYYY-MM-DD，收到：{report_date}")
        try:
            date.fromisoformat(report_date)
        except ValueError:
            raise UserError(f"--date 不是有效日期：{report_date}")
    else:
        if not data["time_min_dt"]:
            raise UserError("导出表里没有可解析的预警时间，请用 --date 指定报表日期。")
        report_date = data["time_min_dt"].date().isoformat()
        print(f"提示：未指定 --date，按数据起始日使用 {report_date}。")

    if not args.false:
        print("提示：未指定任何 --false，所有技能按 0 误报计算。")
    false, specified, all_skills, assumed = parse_false_specs(args.false, data["counts"])
    reasons = parse_reason_specs(args.reason, data["counts"], false)
    wb, records = build_daily_workbook(
        report_date, data, false, specified, all_skills, assumed, reasons,
    )

    outdir = Path(args.outdir).expanduser().resolve() if args.outdir else export.parent
    outdir.mkdir(parents=True, exist_ok=True)
    if args.out:
        out_path = Path(args.out).expanduser().resolve()
    else:
        out_path = outdir / f"准确率日报_{report_date}.xlsx"
    safe_save(wb, out_path)

    ledger_path = None
    if not args.no_ledger:
        if args.ledger:
            ledger_path = Path(args.ledger).expanduser().resolve()
            if ledger_path.is_dir():
                ledger_path = ledger_path / "准确率台账.csv"
        else:
            ledger_path = outdir / "台账" / "准确率台账.csv"
        update_ledger(ledger_path, report_date, records)

    total = data["total"]
    false_total = sum(false.values())
    accurate_total = total - false_total
    print(
        f"日报 {report_date} | 预警 {total} | 误报 {false_total} | "
        f"准确 {accurate_total} | 准确率 {accurate_total / total:.1%}"
    )
    for skill in sorted(data["counts"], key=lambda key: (-data["counts"][key], key)):
        count = data["counts"][skill]
        mis = false[skill]
        extra = f"，原因：{reasons[skill]}" if reasons.get(skill) else ""
        print(
            f"  {skill}: {count} 条，误报 {mis}，"
            f"准确率 {(count - mis) / count:.1%}{extra}"
        )
    missing_reason = [
        f"{skill}（{false[skill]} 条）"
        for skill in data["counts"]
        if false[skill] > 0 and not reasons.get(skill)
    ]
    if missing_reason:
        print("  未填写误报原因：" + "、".join(missing_reason))
    if assumed:
        print("  未提及、按 0 误报：" + "、".join(assumed))
    if data.get("excluded_counts"):
        print(
            "  已排除、不统计："
            + "、".join(
                f"{skill} {count} 条"
                for skill, count in sorted(data["excluded_counts"].items())
            )
        )
    print(f"日报输出：{out_path}")
    if ledger_path:
        print(f"台账：{ledger_path}")
    if not args.no_tidy:
        tidy_reports(outdir)


# ---------------------------------------------------------------------------
# 月报
# ---------------------------------------------------------------------------

def read_ledger(path: Path) -> list[dict]:
    if not path.exists():
        raise UserError(f"台账不存在：{path}。请先跑 daily 生成台账。")
    rows = []
    required = {"报表日期", "预警技能", "预警总数", "误报数"}
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = set(reader.fieldnames or [])
        missing = required - fields
        if missing:
            raise UserError(
                f"台账 {path} 缺少列：{'、'.join(sorted(missing))}。"
            )
        for line_no, item in enumerate(reader, 2):
            try:
                total = int(float(str(item["预警总数"]).strip()))
                mis = int(float(str(item["误报数"]).strip()))
            except (TypeError, ValueError):
                raise UserError(f"台账第 {line_no} 行“预警总数/误报数”不是整数。")
            if total < 0 or mis < 0 or mis > total:
                raise UserError(
                    f"台账第 {line_no} 行数量不合法：总数 {total}，误报 {mis}。"
                )
            report_date = str(item["报表日期"]).strip()
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", report_date):
                raise UserError(f"台账第 {line_no} 行报表日期格式不对：{report_date}")
            try:
                date.fromisoformat(report_date)
            except ValueError:
                raise UserError(f"台账第 {line_no} 行报表日期无效：{report_date}")
            skill = str(item["预警技能"]).strip()
            if not skill:
                raise UserError(f"台账第 {line_no} 行预警技能为空。")
            rows.append({
                "date": report_date,
                "skill": skill,
                "total": total,
                "false": mis,
                "reason": str(item.get("误报原因") or "").strip(),
                "accurate": total - mis,
                "accuracy": (total - mis) / total,
                "source": item.get("数据源", ""),
                "start": item.get("统计开始", ""),
                "end": item.get("统计结束", ""),
                "generated": item.get("生成时间", ""),
            })
    if not rows:
        raise UserError(f"台账为空：{path}")
    return rows


def build_monthly_workbook(month: str, rows: list[dict], today: date):
    try:
        year, mon = [int(part) for part in month.split("-")]
    except ValueError:
        raise UserError(f"--month 格式应为 YYYY-MM，收到：{month}")
    if not 1 <= mon <= 12:
        raise UserError(f"--month 月份不合法：{month}")

    month_start = date(year, mon, 1)
    month_end = date(year, mon, calendar.monthrange(year, mon)[1])
    if today < month_start:
        raise UserError(f"{month} 尚未开始，无法生成月报。")
    cutoff = min(today, month_end)
    days_elapsed = cutoff.day
    day_list = [
        date(year, mon, day).isoformat()
        for day in range(1, days_elapsed + 1)
    ]

    month_rows = [
        item for item in rows
        if item["date"].startswith(month + "-")
        and item["date"] <= cutoff.isoformat()
    ]
    if not month_rows:
        available = sorted({item["date"][:7] for item in rows})
        raise UserError(
            f"台账里没有 {month} 的数据。台账现有月份：{'、'.join(available)}"
        )

    seen = set()
    for item in month_rows:
        key = (item["date"], item["skill"])
        if key in seen:
            raise UserError(
                f"台账中 {item['date']} 的“{item['skill']}”有重复记录，请先清理台账。"
            )
        seen.add(key)

    skills: dict[str, dict] = {}
    days: dict[str, dict] = {}
    for item in month_rows:
        skill_data = skills.setdefault(
            item["skill"],
            {"total": 0, "false": 0, "accurate": 0, "days": set(), "reasons": []},
        )
        skill_data["total"] += item["total"]
        skill_data["false"] += item["false"]
        skill_data["accurate"] += item["accurate"]
        skill_data["days"].add(item["date"])
        reason = item.get("reason", "")
        if reason and reason not in skill_data["reasons"]:
            skill_data["reasons"].append(reason)

        day_data = days.setdefault(
            item["date"], {"total": 0, "false": 0, "accurate": 0},
        )
        day_data["total"] += item["total"]
        day_data["false"] += item["false"]
        day_data["accurate"] += item["accurate"]

    total = sum(item["total"] for item in skills.values())
    false_total = sum(item["false"] for item in skills.values())
    accurate_total = total - false_total
    accuracy_total = accurate_total / total
    days_with_data = len(days)
    coverage = days_with_data / days_elapsed
    avg_daily = total / days_with_data

    skill_order = sorted(skills, key=lambda key: (-skills[key]["total"], key))
    problem = [
        (skill, skills[skill]["accurate"] / skills[skill]["total"])
        for skill in skill_order
        if skills[skill]["accurate"] / skills[skill]["total"] < 0.9
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "月度总览"
    ws.sheet_view.showGridLines = False
    ncols = 8

    add_title(ws, f"{year}年{mon:02d}月 预警准确率月报", ncols)
    add_note(
        ws, 2,
        f"统计区间：{month_start.isoformat()} 至 {cutoff.isoformat()}"
        f"　｜　数据源：准确率台账.csv　｜　有效天数：{days_with_data}/{days_elapsed}",
        ncols,
    )
    add_note(
        ws, 3,
        "口径：准确率 =（预警总数 − 误报数）÷ 预警总数；月准确率按条数加权，"
        "不是每日准确率的简单平均。",
        ncols,
    )

    row = 4
    if coverage < 0.8:
        add_note(
            ws, row,
            f"注意：本月数据完整度仅 {coverage:.1%}（{days_with_data}/{days_elapsed} 天），"
            "以下为阶段性月报，结论仅供参考。",
            ncols, fill=COLOR_WARN_FILL, font_color=COLOR_WARN_FONT, bold=True,
        )
        row += 1
    row += 1

    kpis = [
        ("预警总数", total, "0"),
        ("误报数", false_total, "0"),
        ("准确数", accurate_total, "0"),
        ("月准确率", accuracy_total, "0.0%"),
        ("有效天数", f"{days_with_data}/{days_elapsed}", "@"),
        ("日均预警", avg_daily, "0.0"),
    ]
    for col, (label, value, number_format) in enumerate(kpis, 1):
        label_cell = ws.cell(row, col, label)
        label_cell.font = Font(name=FONT_NAME, size=9, color=COLOR_GREY, bold=True)
        label_cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_LIGHT)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        label_cell.border = thin_border()
        value_cell = ws.cell(row + 1, col, value)
        value_cell.font = Font(name=FONT_NAME, size=15, bold=True, color=COLOR_NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.border = thin_border()
        value_cell.number_format = number_format
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row + 1].height = 34
    row += 3

    add_section(ws, row, "一、按技能汇总", ncols)
    row += 1
    write_header_row(ws, row, [
        "预警技能", "预警总数", "误报数", "误报原因", "准确数", "月准确率", "占总量", "有数据天数",
    ])
    row += 1
    for skill in skill_order:
        item = skills[skill]
        accuracy = item["accurate"] / item["total"]
        low = accuracy < 0.9
        reason_text = (
            "；".join(item["reasons"]) if item["reasons"]
            else ("未填写" if item["false"] > 0 else "—")
        )
        write_row(
            ws, row,
            [
                skill, item["total"], item["false"], reason_text,
                item["accurate"], accuracy, item["total"] / total, len(item["days"]),
            ],
            percent_cols=(6, 7),
            fill=COLOR_BAD_FILL if low else None,
            font_color=COLOR_BAD_FONT if low else None,
            bold=low,
            left_cols=(4,),
        )
        ws.row_dimensions[row].height = 28 if len(reason_text) > 18 else 22
        row += 1
    write_row(
        ws, row,
        ["合计", total, false_total, "", accurate_total, accuracy_total, 1.0, days_with_data],
        percent_cols=(6, 7), bold=True, fill=COLOR_NAVY_LIGHT,
    )
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.border = Border(
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
            top=Side(style="double", color=COLOR_NAVY),
            bottom=Side(style="thin", color=COLOR_BORDER),
        )
    row += 2

    add_section(ws, row, "二、按日明细", ncols)
    row += 1
    write_header_row(ws, row, [
        "日期", "预警总数", "误报数", "准确数", "日准确率", "状态", "备注",
    ])
    row += 1
    for day in day_list:
        if day in days:
            item = days[day]
            accuracy = item["accurate"] / item["total"]
            low = accuracy < 0.9
            write_row(
                ws, row,
                [
                    day, item["total"], item["false"], item["accurate"],
                    accuracy, "有数据", "",
                ],
                percent_cols=(5,),
                fill=COLOR_BAD_FILL if low else None,
                font_color=COLOR_BAD_FONT if low else None,
                bold=low,
            )
        else:
            write_row(
                ws, row,
                [day, "—", "—", "—", "—", "无数据", "台账中没有这一天"],
                fill=COLOR_MISSING_FILL, font_color=COLOR_GREY,
            )
        row += 1
    write_row(
        ws, row,
        ["合计", total, false_total, accurate_total, accuracy_total, "有数据日合计", ""],
        percent_cols=(5,), bold=True, fill=COLOR_NAVY_LIGHT,
    )
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.border = Border(
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
            top=Side(style="double", color=COLOR_NAVY),
            bottom=Side(style="thin", color=COLOR_BORDER),
        )
    row += 2

    missing_days = [day[5:] for day in day_list if day not in days]
    if missing_days:
        shown = "、".join(missing_days[:10])
        if len(missing_days) > 10:
            shown += f" 等 {len(missing_days)} 天"
        missing_text = f"缺失日期：{shown}。"
    else:
        missing_text = "本月每天都有数据。"
    add_note(
        ws, row,
        f"数据完整度：{days_with_data}/{days_elapsed} 天（{coverage:.1%}）。{missing_text}",
        ncols,
        fill=COLOR_WARN_FILL if coverage < 0.8 else None,
        font_color=COLOR_WARN_FONT if coverage < 0.8 else COLOR_GREY,
    )
    row += 1
    if problem:
        add_note(
            ws, row,
            "月准确率低于 90% 的技能：" + "、".join(
                f"{skill} {accuracy:.1%}" for skill, accuracy in problem
            ) + "。",
            ncols, fill=COLOR_BAD_FILL, font_color=COLOR_BAD_FONT, bold=True,
        )
        row += 1
    add_note(
        ws, row,
        "误报数来自每日台账；如日报口径调整，重跑对应日期的 daily 后，月报会自动更新。",
        ncols,
    )
    row += 1
    add_note(ws, row, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ncols)

    widths = {"A": 20, "B": 11, "C": 10, "D": 42, "E": 10, "F": 11, "G": 10, "H": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A9"

    # 技能 x 日期矩阵
    ws2 = wb.create_sheet("技能日矩阵")
    ws2.sheet_view.showGridLines = False
    matrix_cols = 1 + len(day_list) + 1
    add_title(ws2, f"{year}年{mon:02d}月 技能×日期准确率矩阵", matrix_cols, size=13)
    add_note(
        ws2, 2,
        "颜色：红<80%，黄80%–90%，绿≥90%；空白=当天无该技能预警。最后一列为该技能月准确率。",
        matrix_cols,
    )
    matrix_header_row = 4
    headers = ["预警技能"] + [day[5:] for day in day_list] + ["月准确率"]
    write_header_row(ws2, matrix_header_row, headers)

    index = {(item["date"], item["skill"]): item for item in month_rows}
    matrix_row = matrix_header_row + 1
    for skill in skill_order:
        write_row(ws2, matrix_row, [skill], center_first=False)
        for col, day in enumerate(day_list, 2):
            item = index.get((day, skill))
            if item:
                cell = ws2.cell(matrix_row, col, item["accuracy"])
                cell.number_format = "0.0%"
            else:
                cell = ws2.cell(matrix_row, col, None)
            cell.font = Font(name=FONT_NAME, size=9, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border()
        month_accuracy = skills[skill]["accurate"] / skills[skill]["total"]
        cell = ws2.cell(matrix_row, matrix_cols, month_accuracy)
        cell.number_format = "0.0%"
        cell.font = Font(
            name=FONT_NAME, size=9, bold=True,
            color=COLOR_BAD_FONT if month_accuracy < 0.9 else "000000",
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border()
        matrix_row += 1

    # 全部技能行
    write_row(ws2, matrix_row, ["全部技能"], bold=True, fill=COLOR_NAVY_LIGHT)
    for col, day in enumerate(day_list, 2):
        item = days.get(day)
        cell = ws2.cell(matrix_row, col, item["accurate"] / item["total"] if item else None)
        cell.number_format = "0.0%"
        cell.font = Font(name=FONT_NAME, size=9, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border()
        if item:
            cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_LIGHT)
    cell = ws2.cell(matrix_row, matrix_cols, accuracy_total)
    cell.number_format = "0.0%"
    cell.font = Font(name=FONT_NAME, size=9, bold=True, color=COLOR_NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border()
    cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_LIGHT)

    if day_list:
        first_data_row = matrix_header_row + 1
        last_data_row = matrix_row - 1
        last_date_col = 1 + len(day_list)
        rng = (
            f"B{first_data_row}:"
            f"{get_column_letter(last_date_col)}{last_data_row}"
        )
        ws2.conditional_formatting.add(
            rng,
            FormulaRule(
                formula=["AND(ISNUMBER(B5),B5<0.8)"],
                fill=PatternFill("solid", fgColor="F8696B"),
                font=Font(color="9C0006", bold=True),
            ),
        )
        ws2.conditional_formatting.add(
            rng,
            FormulaRule(
                formula=["AND(ISNUMBER(B5),B5>=0.8,B5<0.9)"],
                fill=PatternFill("solid", fgColor="FFEB84"),
                font=Font(color="9C6500"),
            ),
        )
        ws2.conditional_formatting.add(
            rng,
            FormulaRule(
                formula=["AND(ISNUMBER(B5),B5>=0.9)"],
                fill=PatternFill("solid", fgColor="63BE7B"),
                font=Font(color="006100"),
            ),
        )
    ws2.column_dimensions["A"].width = 20
    for col in range(2, matrix_cols):
        ws2.column_dimensions[get_column_letter(col)].width = 6.5
    ws2.column_dimensions[get_column_letter(matrix_cols)].width = 11
    ws2.freeze_panes = "B5"

    # 趋势图
    ws3 = wb.create_sheet("趋势图")
    ws3.sheet_view.showGridLines = False
    add_title(ws3, f"{year}年{mon:02d}月 预警准确率趋势", 12, size=13)
    add_note(ws3, 2, "左表：每日准确率；右表：各技能月准确率。空白日期表示台账无数据。", 12)

    write_header_row(ws3, 4, ["日期", "日准确率"])
    for i, day in enumerate(day_list, 5):
        ws3.cell(i, 1, day[5:]).alignment = Alignment(horizontal="center")
        ws3.cell(i, 1).border = thin_border()
        item = days.get(day)
        cell = ws3.cell(i, 2, item["accurate"] / item["total"] if item else None)
        cell.number_format = "0.0%"
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border()

    for col, text in enumerate(["预警技能", "月准确率"], 4):
        cell = ws3.cell(4, col, text)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=COLOR_NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border()
    for i, skill in enumerate(skill_order, 5):
        cell = ws3.cell(i, 4, skill)
        cell.font = Font(name=FONT_NAME, size=9)
        cell.alignment = Alignment(horizontal="left")
        cell.border = thin_border()
        cell = ws3.cell(i, 5, skills[skill]["accurate"] / skills[skill]["total"])
        cell.number_format = "0.0%"
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border()

    line = LineChart()
    line.title = "每日准确率趋势"
    line.style = 12
    line.y_axis.title = "准确率"
    line.x_axis.title = "日期"
    line.y_axis.numFmt = "0%"
    line.y_axis.scaling.min = 0
    line.y_axis.scaling.max = 1
    line.height = 8.5
    line.width = 19
    data = Reference(ws3, min_col=2, min_row=4, max_row=4 + len(day_list))
    cats = Reference(ws3, min_col=1, min_row=5, max_row=4 + len(day_list))
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.dispBlanksAs = "gap"
    if line.series:
        line.series[0].marker = Marker(symbol="circle", size=6)
    ws3.add_chart(line, "G4")

    bar = BarChart()
    bar.type = "col"
    bar.title = "各技能月准确率"
    bar.style = 10
    bar.y_axis.title = "准确率"
    bar.y_axis.numFmt = "0%"
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = 1
    bar.height = 8.5
    bar.width = 19
    data = Reference(ws3, min_col=5, min_row=4, max_row=4 + len(skill_order))
    cats = Reference(ws3, min_col=4, min_row=5, max_row=4 + len(skill_order))
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.legend = None
    ws3.add_chart(bar, "G23")

    ws3.column_dimensions["A"].width = 10
    ws3.column_dimensions["B"].width = 10
    ws3.column_dimensions["C"].width = 3
    ws3.column_dimensions["D"].width = 20
    ws3.column_dimensions["E"].width = 10

    # 台账
    ws4 = wb.create_sheet("台账")
    ws4.sheet_view.showGridLines = False
    write_header_row(ws4, 1, LEDGER_COLUMNS)
    row4 = 2
    for item in sorted(month_rows, key=lambda x: (x["date"], x["skill"])):
        write_row(
            ws4, row4,
            [
                item["date"], item["skill"], item["total"], item["false"],
                item["reason"], item["accurate"], item["accuracy"], item["source"],
                item["start"], item["end"], item["generated"],
            ],
            percent_cols=(7,),
            left_cols=(5,),
        )
        row4 += 1
    widths4 = {
        "A": 12, "B": 18, "C": 10, "D": 10, "E": 40,
        "F": 10, "G": 10, "H": 26, "I": 20, "J": 20, "K": 20,
    }
    for col, width in widths4.items():
        ws4.column_dimensions[col].width = width
    ws4.freeze_panes = "A2"

    summary = {
        "total": total,
        "false": false_total,
        "accurate": accurate_total,
        "accuracy": accuracy_total,
        "days_with_data": days_with_data,
        "days_elapsed": days_elapsed,
        "coverage": coverage,
        "problem": problem,
    }
    return wb, summary


def default_monthly_outdir(ledger_path: Path) -> Path:
    """台账放在 台账/ 子目录时，月报输出到它的上一级。"""
    if ledger_path.parent.name in {"台账", "台账数据", "data"}:
        return ledger_path.parent.parent
    return ledger_path.parent


def cmd_monthly(args) -> None:
    month = args.month.strip()
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise UserError(f"--month 格式应为 YYYY-MM，收到：{month}")

    if args.today:
        try:
            today = date.fromisoformat(args.today)
        except ValueError:
            raise UserError(f"--today 不是有效日期：{args.today}")
    else:
        today = date.today()

    ledger_path = Path(args.ledger).expanduser().resolve()
    if ledger_path.is_dir():
        ledger_path = ledger_path / "准确率台账.csv"
    rows = read_ledger(ledger_path)
    wb, summary = build_monthly_workbook(month, rows, today)

    if args.out:
        out_path = Path(args.out).expanduser().resolve()
    else:
        out_path = default_monthly_outdir(ledger_path) / f"准确率月报_{month}.xlsx"
    safe_save(wb, out_path)

    print(
        f"月报 {month} | 预警 {summary['total']} | 误报 {summary['false']} | "
        f"准确 {summary['accurate']} | 准确率 {summary['accuracy']:.1%} | "
        f"有效天数 {summary['days_with_data']}/{summary['days_elapsed']}"
    )
    if summary["problem"]:
        print(
            "  低于 90%：" + "、".join(
                f"{skill} {accuracy:.1%}"
                for skill, accuracy in summary["problem"]
            )
        )
    print(f"月报输出：{out_path}")
    if not args.no_tidy:
        tidy_reports(out_path.parent)


def cmd_tidy(args) -> None:
    outdir = (
        Path(args.outdir).expanduser().resolve()
        if args.outdir else Path.cwd()
    )
    if not outdir.is_dir():
        raise UserError(f"工作目录不存在：{outdir}")
    moved, skipped, problems = tidy_reports(outdir)
    if not moved and not skipped and not problems:
        print("没有需要整理的旧报告。")


def cmd_archive(args) -> None:
    month = args.month.strip()
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise UserError(f"--month 格式应为 YYYY-MM，收到：{month}")
    year, mon = [int(part) for part in month.split("-")]
    if not 1 <= mon <= 12:
        raise UserError(f"--month 月份不合法：{month}")
    try:
        date(year, mon, 1)
    except ValueError:
        raise UserError(f"--month 不是有效月份：{month}")

    today_month = date.today().strftime("%Y-%m")
    if month > today_month:
        raise UserError(f"{month} 还没到，不能归档。")
    if month == today_month and not args.force:
        raise UserError(
            f"{month} 还没结束。确认要归档当前月份，请加 --force。"
        )

    outdir = (
        Path(args.outdir).expanduser().resolve()
        if args.outdir else Path.cwd()
    )
    if not outdir.is_dir():
        raise UserError(f"工作目录不存在：{outdir}")

    archive_root = outdir / args.archive_root if args.archive_root else outdir
    archive_dir = archive_root / month
    if archive_dir.exists() and not archive_dir.is_dir():
        raise UserError(f"归档路径已存在且不是文件夹：{archive_dir}")

    daily_files = sorted(outdir.glob(f"准确率日报_{month}-*.xlsx"))
    monthly_file = outdir / f"准确率月报_{month}.xlsx"
    pairs = [
        (src, archive_dir / "日报" / src.name)
        for src in daily_files
    ]
    if monthly_file.exists():
        pairs.append((monthly_file, archive_dir / monthly_file.name))
    else:
        print(
            f"提示：没有找到 {monthly_file.name}，本次只归档日报；"
            "月报生成后再跑一次 archive 即可。"
        )
    if not pairs:
        print(f"没有找到 {month} 的日报或月报，未执行归档。")
        return

    conflicts = []
    for src, dest in pairs:
        if dest.exists() and not (
            dest.is_file() and filecmp.cmp(src, dest, shallow=False)
        ):
            conflicts.append(dest)
    if conflicts:
        raise UserError(
            "归档目标已存在且内容不同，请先检查：\n"
            + "\n".join(str(path) for path in conflicts)
        )

    if args.dry_run:
        print(f"预览：将归档 {len(pairs)} 个文件到 {archive_dir}")
        for src, dest in pairs:
            print(f"  {src.name} -> {dest}")
        return

    archive_dir.mkdir(parents=True, exist_ok=True)
    moved = 0
    skipped = 0
    problems = []
    for src, dest in pairs:
        result = _move_report(src, dest)
        if result == "moved":
            moved += 1
        elif result == "skipped":
            skipped += 1
        else:
            problems.append(src.name)
    print(
        f"归档完成：移动 {moved} 个，跳过已存在 {skipped} 个 → {archive_dir}"
    )
    if problems:
        print("  未归档：" + "、".join(problems))


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="预警准确率日报 / 月报生成器",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "示例：\n"
            "  python accuracy_report.py daily 26日-27日早9点.xlsx --date 2026-10-26 "
            "--false \"夜间区域入侵=3\" --false \"区域入侵=1\" --false \"烟雾识别=all\" "
            "--reason \"夜间区域入侵=算法误识别：把门识别成人\" "
            "--reason \"区域入侵=算法误识别：022镜头把围栏识别成人\"\n"
            "  python accuracy_report.py monthly --month 2026-10 "
            "--ledger 台账/准确率台账.csv\n"
            "  python accuracy_report.py archive --month 2026-10 "
            "--outdir <工作目录>"
        ),
    )
    parser.add_argument("--debug", action="store_true", help="出错时打印完整堆栈")
    sub = parser.add_subparsers(dest="command", required=True)

    daily = sub.add_parser("daily", help="生成某一天的准确率日报并更新台账")
    daily.add_argument("export", help="预警平台导出的 xlsx 文件")
    daily.add_argument("--date", help="报表日期 YYYY-MM-DD；不填则取数据中最早一条的日期")
    daily.add_argument(
        "--false", action="append", default=[], metavar="技能=数量|all",
        help="某技能的误报条数，可重复；all 表示该技能全部误报",
    )
    daily.add_argument(
        "--reason", action="append", default=[], metavar="技能=原因",
        help="某技能的误报原因，可重复；通常与 --false 配套，未填写的显示“未填写”",
    )
    daily.add_argument(
        "--exclude", action="append", default=[], metavar="技能",
        help="不统计的技能，可重复；其预警条数不计入总数和准确率",
    )
    daily.add_argument("--sheet", help="工作表名，默认自动查找包含“预警技能”的表")
    daily.add_argument(
        "--list-skills", action="store_true",
        help="只列出导出表中的技能和条数，不生成日报",
    )
    daily.add_argument("--outdir", help="输出目录，默认与导出文件同目录")
    daily.add_argument("--out", help="日报 xlsx 的完整输出路径")
    daily.add_argument("--ledger", help="台账 CSV 路径，默认输出目录/台账/准确率台账.csv")
    daily.add_argument("--no-ledger", action="store_true", help="不更新台账")
    daily.add_argument(
        "--no-tidy", action="store_true", help="不自动把旧日报/月报收进 归档/"
    )
    daily.set_defaults(func=cmd_daily)

    monthly = sub.add_parser("monthly", help="按月份汇总台账，生成月报")
    monthly.add_argument("--month", required=True, help="月份 YYYY-MM")
    monthly.add_argument("--ledger", required=True, help="准确率台账.csv 路径")
    monthly.add_argument("--out", help="月报 xlsx 的完整输出路径")
    monthly.add_argument("--today", help="调试用：把“今天”固定为 YYYY-MM-DD")
    monthly.add_argument(
        "--no-tidy", action="store_true", help="不自动把旧日报/月报收进 归档/"
    )
    monthly.set_defaults(func=cmd_monthly)

    tidy = sub.add_parser("tidy", help="手动把旧日报/月报整理到 归档/YYYY-MM/")
    tidy.add_argument("--outdir", help="工作目录，默认当前目录")
    tidy.set_defaults(func=cmd_tidy)

    archive = sub.add_parser(
        "archive", help="把某个月的日报和月报移动到 归档/YYYY-MM/"
    )
    archive.add_argument("--month", required=True, help="月份 YYYY-MM")
    archive.add_argument("--outdir", help="工作目录，默认当前目录")
    archive.add_argument(
        "--archive-root", default="归档",
        help="归档根目录名，默认 归档；设为空字符串则直接建 YYYY-MM/",
    )
    archive.add_argument(
        "--dry-run", action="store_true", help="只预览要移动的文件，不实际移动"
    )
    archive.add_argument(
        "--force", action="store_true", help="允许归档当前月份"
    )
    archive.set_defaults(func=cmd_archive)

    return parser


def main(argv=None) -> int:
    _init_stdout()
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        args.func(args)
    except UserError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1
    except PermissionError as exc:
        print(f"错误：文件被占用或无权限：{exc}", file=sys.stderr)
        return 1
    except Exception as exc:  # pragma: no cover
        if getattr(args, "debug", False):
            traceback.print_exc()
        else:
            print(f"错误：{exc}", file=sys.stderr)
            print("可加 --debug 查看完整堆栈。", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
