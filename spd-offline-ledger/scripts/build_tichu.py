# -*- coding: utf-8 -*-
"""build_tichu.py — 生成"剔除后表"（每管理处一张，按运维处归档）+ 总表
用法: python build_tichu.py <工作目录>
输入: 目录下 北京.csv/河北.csv/天津.csv (GBK; 二级分组=管理处)
规则: 剔除(分组含隐藏/二级为隐藏目录·隐藏点位·未分组/分组均空/掉线>8640);
      同一摄像机(管理处+名称+IP)去重取最大值; 按掉线时长正序; 批量(同值>=3且>=500)标黄.
注: 2026-09 起口径变为 P=Q+R+S+T+U, 周界计入 S, 因此"分组含周界"不再剔除, 周界行保留在剔除后表里。
加固: 单个管理处写入失败不中断其余(打印错误); 清洗非法字符; 名称以=开头防公式.
"""
import io, os, sys, glob, re, traceback, collections

BASE = sys.argv[1] if len(sys.argv) > 1 else r"."
OUT  = os.path.join(BASE, "剔除后表_按运维处")
os.makedirs(OUT, exist_ok=True)
SIX = 8640

_ILL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f\ufeff\ufffe\uffff]')
def clean(v):
    if isinstance(v, str):
        v = _ILL.sub('', v)
        if v.startswith('='):
            v = "'" + v
        return v
    return v

def parse(fn):
    with io.open(fn, 'r', encoding='gbk') as fh:
        lines = fh.readlines()
    interval = lines[0].strip().replace('统计区间: ', '')
    hdr = next(i for i, l in enumerate(lines) if l.strip().startswith('序号,一级分组'))
    header = [c.strip() for c in lines[hdr].strip().split(',')]
    idx = {c: header.index(c) for c in header if c}
    def gc(p, n):
        return clean(p[idx[n]]) if n in idx and idx[n] < len(p) else ''
    rows = []
    for l in lines[hdr+1:]:
        p = l.strip().split(',')
        if p and p[0].strip().isdigit() and len(p) >= 8:
            rows.append(dict(f=os.path.basename(fn), org=gc(p, '一级分组'), org2=gc(p, '二级分组'),
                             g3=gc(p, '三级分组'), g4=gc(p, '四级分组'), g5=gc(p, '五级分组'), g6=gc(p, '六级分组'),
                             name=gc(p, '镜头名称'), ip=gc(p, '设备IP地址'),
                             dur=int(gc(p, '掉线时长') or 0), cnt=int(gc(p, '掉线次数') or 0)))
    return interval, rows

# 分公司文件 + 公司级统计
intervals = {}; branch_sum = {}; allrows = []
for fn in ['北京.csv', '天津.csv', '河北.csv']:
    fp = os.path.join(BASE, fn)
    if not os.path.exists(fp):
        continue
    iv, rows = parse(fp)
    intervals[fn] = iv; allrows.extend(rows)
    with io.open(fp, 'r', encoding='gbk') as fh:
        lines = fh.readlines()
    hdr = next(i for i, l in enumerate(lines) if l.strip().startswith('序号,一级分组'))
    for l in lines[hdr+1:]:
        p = l.strip().split(',')
        if p and p[0] == '全局' and len(p) >= 5:
            branch_sum[fn] = dict(总镜头=p[1], 在线=p[2], 掉线=p[3], 在线率=p[4])
        if p and p[0] == '' and len(p) >= 5 and p[1].strip().isdigit():
            branch_sum.setdefault(fn, {}).update(总在线时长=p[1], 掉线时长=p[3], 时长率=p[4])

# 合并去重(管理处+名称+IP 取最大)
merged = collections.defaultdict(dict)
for r in allrows:
    k = (r['org2'], r['name'], r['ip'])
    if k not in merged[r['org2']] or r['dur'] > merged[r['org2']][k]['dur']:
        merged[r['org2']][k] = r

def remove_reason(r):
    if any('隐藏' in r[k] for k in ('org', 'org2', 'g3', 'g4', 'g5', 'g6')): return '分组含隐藏'
    if r['org2'] in ('隐藏目录', '隐藏点位', '未分组'): return '二级分组为隐藏/未分组'
    if (r['org2'] == '' and r['g3'] == '' and r['g4'] == '' and r['g5'] == '' and r['g6'] == ''): return '分组均为空'
    if r['dur'] > SIX: return '掉线时长超6天'
    return None

kept_by_org2 = {}; removed_by_org2 = {}; removed_log = collections.Counter()
for org2, camdict in merged.items():
    kept = []; removed = []
    for k, r in camdict.items():
        t = remove_reason(r)
        if t:
            removed_log[t] += 1; removed.append((r, t))
        else:
            kept.append(r)
    kept_by_org2[org2] = kept; removed_by_org2[org2] = removed

# 保定与顺平合并在同一张表里(二级分组都叫保定管理处): 按四级/五级分组含"顺平"拆成两个管理处,
# 让掉线时长/剔除量/在线率等各项计算各自独立。
if '保定管理处' in kept_by_org2:
    def _is_sp(r):
        return '顺平' in (str(r['g4'] or '') + str(r['g5'] or ''))
    kept_by_org2['顺平管理处'] = [r for r in kept_by_org2['保定管理处'] if _is_sp(r)]
    kept_by_org2['保定管理处'] = [r for r in kept_by_org2['保定管理处'] if not _is_sp(r)]
    _rem = removed_by_org2.get('保定管理处', [])
    removed_by_org2['顺平管理处'] = [x for x in _rem if _is_sp(x[0])]
    removed_by_org2['保定管理处'] = [x for x in _rem if not _is_sp(x[0])]

all_kept = [r for v in kept_by_org2.values() for r in v]
hist = collections.Counter(r['dur'] for r in all_kept)
def batch_flag(r):
    c = hist[r['dur']]
    return f"批量掉线(同{r['dur']}分钟x{c}行)" if (r['dur'] >= 500 and c >= 3) else ''

real_org2s = [o for o in kept_by_org2 if o not in ('隐藏目录', '隐藏点位', '未分组')]
for o in real_org2s:
    kept_by_org2[o].sort(key=lambda r: (r['dur'], r['name']))

YW_MAP = {
    '邯郸运维处': ['磁县管理处', '邯郸管理处', '永年管理处'],
    '邢台运维处': ['沙河管理处', '邢台管理处', '临城管理处'],
    '石家庄运维处': ['高邑元氏管理处', '石家庄管理处', '新乐管理处', '河北分公司本部'],
    '保定运维处': ['定州管理处', '唐县管理处', '顺平管理处', '保定管理处'],
    '天津运维处': ['西黑山管理处', '徐水管理处', '容雄管理处', '霸州管理处', '天津管理处', '天津分公司本部'],
    '北京运维处': ['易县管理处', '涞涿管理处', '惠南庄管理处'],
}
NOCOUNT = {'容雄管理处'}
def org2_yw(o):
    for yw, lst in YW_MAP.items():
        if o in lst: return yw
    return '未分组'

print("剔除原因:", dict(removed_log))
print(f"\n{'管理处':<12}{'保留行':>5}{'剔除行':>5}")

sys.path.insert(0, r"C:\Users\92586\.zcode\cli\plugins\cache\zcode-plugins-official\document-skills\0.1.2\skills\xlsx\templates")
try:
    from base import FONT_NAME, PRIMARY, SECONDARY, NEUTRAL_900, NEUTRAL_100, NEUTRAL_0, NEUTRAL_200, NEUTRAL_600, ACCENT_WARNING, HEADER_BOLD
except Exception:
    FONT_NAME = 'Microsoft YaHei'; PRIMARY = '1B2A4A'; SECONDARY = 'D6E4F0'
    NEUTRAL_900 = '37352F'; NEUTRAL_100 = 'F7F7F5'; NEUTRAL_0 = 'FFFFFF'; NEUTRAL_200 = 'E9E9E8'
    NEUTRAL_600 = '8C8A84'; ACCENT_WARNING = 'D4820A'; HEADER_BOLD = False
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
thin = Side(style='thin', color=NEUTRAL_200); med = Side(style='medium', color=NEUTRAL_200)
hdr_fill = PatternFill('solid', fgColor=PRIMARY); tot_fill = PatternFill('solid', fgColor=SECONDARY); warn_color = 'FEF9E7'
headers = ['序号', '一级分组', '二级分组', '三级分组', '四级分组', '五级分组', '六级分组', '镜头名称', '设备IP地址', '掉线时长', '掉线次数', '批量掉线标记']

failed = []
for o in real_org2s:
    kept = kept_by_org2[o]
    if not kept:
        continue
    try:
        wb = Workbook(); ws = wb.active; ws.title = clean(o)[:28]
        br = kept[0]['org']; fn = kept[0]['f']
        ws['A1'] = clean(f"统计区间: {intervals.get(fn,'')}")
        ws['A1'].font = Font(name=FONT_NAME, size=9, color=NEUTRAL_900)
        bs = branch_sum.get(fn, {})
        ws['A3'] = clean(f"{o} 剔除后表 | {br}总镜头{bs.get('总镜头','?')}台 在线率{bs.get('在线率','?')} | 已剔除: 分组含隐藏/隐藏点位/超6天; 周界行保留(计入新表S); 黄色=批量掉线(待筛选)")
        ws['A3'].font = Font(name=FONT_NAME, size=9, color=NEUTRAL_600)
        hrow = 4
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=hrow, column=ci, value=h); c.fill = hdr_fill
            c.font = Font(name=FONT_NAME, size=10, bold=HEADER_BOLD, color='FFFFFF')
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = Border(bottom=thin)
        ws.row_dimensions[hrow].height = 24
        r = hrow + 1
        for i, rec in enumerate(kept, 1):
            fl = batch_flag(rec)
            vals = [i, clean(rec['org']), clean(rec['org2']), clean(rec['g3']), clean(rec['g4']),
                    clean(rec['g5']), clean(rec['g6']), clean(rec['name']), clean(rec['ip']),
                    rec['dur'], rec['cnt'], fl]
            for ci, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=ci, value=v)
                c.font = Font(name=FONT_NAME, size=9, color=NEUTRAL_900)
                c.alignment = Alignment(horizontal='right' if ci in (1, 9, 10, 11) else 'left', vertical='center')
                c.fill = PatternFill('solid', fgColor=warn_color if fl else (NEUTRAL_0 if i % 2 else NEUTRAL_100))
                if fl: c.font = Font(name=FONT_NAME, size=9, color=ACCENT_WARNING)
            ws.row_dimensions[r].height = 16
            r += 1
        r += 1
        off = sum(x['dur'] for x in kept); rem = removed_by_org2[o]
        for label, val in [('掉线镜头数量(剔除后)', len(kept)), ('掉线镜头时长(分钟)', off),
                           ('剔除计划性掉线镜头数量', len(rem)), ('剔除计划性掉线时长(分钟)', sum(x['dur'] for x, c in rem))]:
            c1 = ws.cell(row=r, column=1, value=label); c2 = ws.cell(row=r, column=2, value=val)
            for cc in (c1, c2):
                cc.font = Font(name=FONT_NAME, size=10, color=NEUTRAL_900); cc.fill = tot_fill
            ws.row_dimensions[r].height = 18; r += 1
        note = (f"公司级口径: {br}总镜头{bs.get('总镜头','?')}台 在线率{bs.get('在线率','?')} "
                f"镜头在线时长率{bs.get('时长率','?')}; 本表未提供分管理处镜头总数. 批量{len([1 for x in kept if batch_flag(x)])}行.")
        c = ws.cell(row=r, column=1, value=clean(note)); c.font = Font(name=FONT_NAME, size=8, color=NEUTRAL_600)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12); ws.row_dimensions[r].height = 26
        for ci, w in enumerate([5, 10, 10, 12, 12, 12, 12, 30, 15, 9, 8, 22], 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.sheet_view.showGridLines = False
        yw = '未计入_容雄' if o in NOCOUNT else org2_yw(o)
        os.makedirs(os.path.join(OUT, yw), exist_ok=True)
        wb.save(os.path.join(OUT, yw, f"{o}_剔除后表.xlsx"))
        print(f"  OK {o:<12}{len(kept):>5}{len(rem):>5}")
    except Exception as e:
        failed.append(o)
        print(f"  [ERROR] {o}: {e}")
        traceback.print_exc()

print(f"\n已生成 {len([o for o in real_org2s if kept_by_org2[o]]) - len(failed)} 个管理处剔除后表 -> {OUT}")
if failed:
    print(f"!! 失败的: {failed}")

# ===== 总表 =====
try:
    wb2 = Workbook(); ws = wb2.active; ws.title = '总表'
    ws['A1'] = "摄像机掉线时长统计总表"; ws.merge_cells('A1:J1')
    ws['A1'].font = Font(name=FONT_NAME, size=14, bold=HEADER_BOLD, color=PRIMARY); ws.row_dimensions[1].height = 30
    hdr = ['分公司', '总镜头数量', '在线镜头数量', '掉线镜头数量', '在线率(镜头数)', '总镜头在线时长', '掉线镜头时长(分钟)', '镜头在线时长率', '剔除计划性掉线镜头数', '剔除计划性掉线时长(分钟)']
    for ci, h in enumerate(hdr, 1):
        c = ws.cell(row=2, column=ci, value=h); c.fill = hdr_fill
        c.font = Font(name=FONT_NAME, size=10, bold=HEADER_BOLD, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = Border(bottom=thin)
    ws.row_dimensions[2].height = 28
    r = 3
    for key in ['北京', '天津', '河北']:
        fn = f'{key}.csv'; bs = branch_sum.get(fn, {})
        cnt = [o for o in real_org2s if o not in NOCOUNT and any(x['f'] == fn for x in kept_by_org2[o])]
        rem_n = sum(len(removed_by_org2[o]) for o in cnt); rem_d = sum(sum(x['dur'] for x, c in removed_by_org2[o]) for o in cnt)
        vals = [f'{key}分公司', bs.get('总镜头', ''), bs.get('在线', ''), bs.get('掉线', ''), bs.get('在线率', ''),
                bs.get('总在线时长', ''), bs.get('掉线时长', ''), bs.get('时长率', ''), rem_n, rem_d]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=ci, value=v); c.font = Font(name=FONT_NAME, size=10, color=NEUTRAL_900)
            c.alignment = Alignment(horizontal='right' if ci in (2, 3, 4, 6, 7, 9, 10) else 'left', vertical='center')
        ws.row_dimensions[r].height = 18; r += 1
    r += 1
    ws.cell(row=r, column=1, value='各管理处明细').font = Font(name=FONT_NAME, size=11, bold=HEADER_BOLD, color=PRIMARY); r += 1
    for ci, h in enumerate(['管理处', '分公司', '保留行数', '掉线时长(分钟)', '剔除行数', '剔除时长(分钟)'], 1):
        c = ws.cell(row=r, column=ci, value=h); c.fill = tot_fill
        c.font = Font(name=FONT_NAME, size=10, bold=HEADER_BOLD, color=PRIMARY); c.border = Border(top=med)
    r += 1
    for o in sorted(real_org2s, key=lambda o: -sum(x['dur'] for x in kept_by_org2[o])):
        kept = kept_by_org2[o]; rem = removed_by_org2[o]
        vals = [o, kept[0]['org'] if kept else '?', len(kept), sum(x['dur'] for x in kept), len(rem), sum(x['dur'] for x, c in rem)]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=ci, value=v); c.font = Font(name=FONT_NAME, size=9, color=NEUTRAL_900)
            c.alignment = Alignment(horizontal='right' if ci in (3, 4, 5, 6) else 'left', vertical='center')
        ws.row_dimensions[r].height = 16; r += 1
    for ci, w in enumerate([12, 8, 10, 14, 10, 14], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.sheet_view.showGridLines = False
    wb2.save(os.path.join(BASE, '总表.xlsx'))
    print("总表.xlsx 已生成")
except Exception as e:
    print(f"[ERROR] 总表: {e}")
    traceback.print_exc()
